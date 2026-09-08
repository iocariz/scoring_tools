"""Tests for the How-We-Got-Here sheet and its matplotlib renderings."""

import openpyxl
import pandas as pd
import pytest

from src.consolidation import _collect_frontier_specs, _write_sheet_how_we_got_here
from src.consolidation_charts import render_frontier_small_multiples, render_production_bridge

TOTALS = {
    "actual_production": 336_800_000.0,
    "swap_out_production": 55_400_000.0,
    "swap_in_production": 22_600_000.0,
    "optimum_production": 304_000_000.0,
    "actual_risk_pct": 1.92,
    "optimum_risk_pct": 1.17,
}


def test_render_production_bridge(tmp_path):
    out = render_production_bridge(TOTALS, tmp_path / "bridge.png")
    assert out is not None and out.exists() and out.stat().st_size > 0


def test_render_production_bridge_missing_fields_returns_none(tmp_path):
    assert render_production_bridge({"actual_production": 1.0}, tmp_path / "bridge.png") is None


def test_render_frontier_small_multiples(tmp_path):
    specs = [
        {
            "name": "seg_a",
            "risk": [0.5, 1.0, 2.0, 3.0],
            "production": [1e6, 2e6, 3e6, 3.5e6],
            "chosen_risk": 1.0,
            "chosen_production": 2e6,
            "target": 1.1,
        },
        # no chosen point / no target — panel still renders
        {"name": "seg_b", "risk": [1.0, 2.0], "production": [1e6, 2e6], "chosen_risk": None, "chosen_production": None},
    ]
    out = render_frontier_small_multiples(specs, tmp_path / "frontiers.png")
    assert out is not None and out.exists() and out.stat().st_size > 0


def test_render_frontier_filters_degenerate_specs(tmp_path):
    # fewer than 2 frontier points cannot draw a curve; all-degenerate -> no image
    specs = [{"name": "seg", "risk": [1.0], "production": [1e6]}]
    assert render_frontier_small_multiples(specs, tmp_path / "frontiers.png") is None


def _make_segment_artifacts(output_base, seg="seg_a", target=1.1):
    data_dir = output_base / seg / "data"
    data_dir.mkdir(parents=True)
    pd.DataFrame({"b2_ever_h6": [0.5, 1.0, 2.0], "oa_amt_h0": [1e6, 2e6, 3e6]}).to_csv(
        data_dir / "efficient_frontier_base.csv", index=False
    )
    pd.DataFrame({"b2_ever_h6": [1.0], "oa_amt_h0": [2e6]}).to_csv(data_dir / "optimal_solution_base.csv", index=False)
    pd.DataFrame({"cell": range(7)}).to_csv(data_dir / "data_summary_desagregado_base.csv", index=False)
    (output_base / seg / "config_segment.toml").write_text(f"[preprocessing]\noptimum_risk = {target}\n")


def test_collect_frontier_specs(tmp_path):
    _make_segment_artifacts(tmp_path)
    specs, n_cells, n_frontier = _collect_frontier_specs(tmp_path, {"seg_a": {}, "missing_seg": {}}, "_base")
    assert len(specs) == 1  # missing_seg has no artifacts -> dropped, not crashed
    assert n_cells == 7 and n_frontier == 3
    s = specs[0]
    assert s["target"] == pytest.approx(1.1)
    assert s["chosen_risk"] == pytest.approx(1.0)
    assert s["chosen_production"] == pytest.approx(2e6)


def _consolidated_df():
    return pd.DataFrame([{"group": "TOTAL", "period": "main", "scenario": "base", "total_demand": 5e8, **TOTALS}])


def test_write_sheet_how_we_got_here(tmp_path):
    _make_segment_artifacts(tmp_path)
    wb = openpyxl.Workbook()
    _write_sheet_how_we_got_here(wb, tmp_path, {"seg_a": {}}, _consolidated_df(), "_base")
    ws = wb["How We Got Here"]
    assert len(ws._images) == 2  # bridge + frontiers embedded
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Start from every application" in text
    assert "7 score cells" in text  # run's actual numbers flow into the step band
    assert "3 candidate accept/reject policies" in text
    assert (tmp_path / "consolidated_charts" / "production_bridge_base.png").exists()
    assert (tmp_path / "consolidated_charts" / "frontiers_base.png").exists()


def test_write_sheet_degrades_without_artifacts(tmp_path):
    """No segment artifacts and no TOTAL row: the sheet still writes, with note cells."""
    wb = openpyxl.Workbook()
    _write_sheet_how_we_got_here(
        wb, tmp_path, {"seg_a": {}}, pd.DataFrame(columns=["group", "period", "scenario"]), "_base"
    )
    ws = wb["How We Got Here"]
    assert len(ws._images) == 0
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Bridge chart" in text and "Frontier charts" in text  # fallback notes present
