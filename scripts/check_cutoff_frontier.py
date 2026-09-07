"""Verify the Cutoff Grids frontier borders are continuous staircases.

A QA helper for the consolidated workbook (companion to
``preview_exec_summary_kpis.py``): it parses every A/R/— acceptance grid on the
"Cutoff Grids" sheet, reconstructs the thick-navy frontier edges cell by cell,
and checks the invariants the grid writer is supposed to guarantee (#215/#219/#220):

1. **No gaps** — flood-filling from the accepted (``A``) cells without crossing a
   frontier edge must never reach a rejected (``R``) cell.
2. **Continuity** — frontier edges form closed chains: every interior lattice
   vertex touches an even number of frontier edges (a dangling edge = a visible
   break in the staircase). Chains may terminate on the grid boundary.
3. **Staircase shape** — the accept side enclosed by the frontier is a monotone
   region: per row one contiguous run anchored at a consistent grid edge, with
   run lengths monotone across rows. Grey (``—``) cells inside the accept region
   count as accept-side (border drawing imputes them), so holes fail loudly.

Usage:
    uv run python scripts/check_cutoff_frontier.py
    uv run python scripts/check_cutoff_frontier.py path/to.xlsx --sheet "Cutoff Grids"

Exit code 0 when every grid passes, 1 otherwise.
"""

from __future__ import annotations

import argparse
import sys

import numpy as np
import openpyxl

# _SIDE_FRONTIER in src/consolidation.py is thick navy; the normal cell
# separator (_SIDE_GRID / _BORDER_GRID) is medium white.
FRONTIER_STYLE = "thick"
CORNER_SEP = " \\ "  # corner label is "<row_var> \\ <col_var>"


def _is_corner(value) -> bool:
    return isinstance(value, str) and CORNER_SEP in value


def find_grids(ws) -> list[tuple[int, int, str, int, int]]:
    """Locate every pivot grid on the sheet.

    Returns (corner_row, corner_col, label, nrows, ncols) per grid; the corner
    cell holds the "row_var \\ col_var" label with headers to its right and below.
    """
    grids = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if not _is_corner(v):
                continue
            ncols = 0
            while (h := ws.cell(row=r, column=c + 1 + ncols).value) is not None and not _is_corner(h):
                ncols += 1
            nrows = 0
            while (h := ws.cell(row=r + 1 + nrows, column=c).value) is not None and not _is_corner(h):
                nrows += 1
            if nrows and ncols:
                grids.append((r, c, v, nrows, ncols))
    return grids


def read_grid(ws, r0: int, c0: int, nrows: int, ncols: int):
    """Read cell values and per-side frontier flags for the grid at corner (r0, c0).

    Returns (vals, frontier) where vals is an object array of "A"/"R"/"—" and
    frontier is a dict of four boolean arrays keyed "top"/"bottom"/"left"/"right".
    """
    vals = np.full((nrows, ncols), "", dtype=object)
    frontier = {side: np.zeros((nrows, ncols), bool) for side in ("top", "bottom", "left", "right")}
    for i in range(nrows):
        for j in range(ncols):
            cell = ws.cell(row=r0 + 1 + i, column=c0 + 1 + j)
            vals[i, j] = cell.value
            b = cell.border
            for side in frontier:
                s = getattr(b, side)
                frontier[side][i, j] = s is not None and s.style == FRONTIER_STYLE
    return vals, frontier


def _accept_side(vals, frontier) -> np.ndarray:
    """Flood-fill from the A cells without crossing frontier edges → accept-side mask."""
    nrows, ncols = vals.shape
    acc = np.zeros((nrows, ncols), bool)
    stack = list(zip(*np.where(vals == "A"), strict=True))
    seen = set(stack)
    while stack:
        i, j = stack.pop()
        acc[i, j] = True
        for di, dj, blocked in (
            (-1, 0, frontier["top"][i, j]),
            (1, 0, frontier["bottom"][i, j]),
            (0, -1, frontier["left"][i, j]),
            (0, 1, frontier["right"][i, j]),
        ):
            ni, nj = i + di, j + dj
            if 0 <= ni < nrows and 0 <= nj < ncols and (ni, nj) not in seen and not blocked:
                seen.add((ni, nj))
                stack.append((ni, nj))
    return acc


def check_grid(vals, frontier) -> list[str]:
    """Run the gap / continuity / staircase checks on one grid; returns problems."""
    nrows, ncols = vals.shape
    has_frontier = any(f.any() for f in frontier.values())
    if not (vals == "A").any() or not has_frontier:
        # nothing accepted (or nothing drawn, e.g. all-reject) → no boundary to validate
        return []
    problems: list[str] = []
    acc = _accept_side(vals, frontier)

    # 1. gaps: the fill must never leak onto a rejected cell
    leaked = [(i, j) for i, j in zip(*np.where(acc), strict=True) if vals[i, j] == "R"]
    if leaked:
        problems.append(f"frontier GAP: flood-fill from A reaches R at {leaked[:5]}")

    # 2. continuity: frontier edges as segments on the (nrows+1)x(ncols+1) vertex lattice
    edges = set()
    for i in range(nrows):
        for j in range(ncols):
            if frontier["top"][i, j]:
                edges.add(((i, j), (i, j + 1)))
            if frontier["bottom"][i, j]:
                edges.add(((i + 1, j), (i + 1, j + 1)))
            if frontier["left"][i, j]:
                edges.add(((i, j), (i + 1, j)))
            if frontier["right"][i, j]:
                edges.add(((i, j + 1), (i + 1, j + 1)))
    deg: dict[tuple[int, int], int] = {}
    for a, b in edges:
        deg[a] = deg.get(a, 0) + 1
        deg[b] = deg.get(b, 0) + 1
    dangling = [v for v, d in deg.items() if d % 2 and v[0] not in (0, nrows) and v[1] not in (0, ncols)]
    if dangling:
        problems.append(f"frontier DISCONTINUOUS: odd-degree interior vertices {dangling[:5]}")

    # 3. staircase shape of the accept side
    anchor = None  # which grid edge the accept runs hug: "left" or "right"
    lens = []
    for i in range(nrows):
        idx = np.where(acc[i])[0]
        lens.append(int(idx.size))
        if idx.size == 0:
            continue
        if idx.size != idx[-1] - idx[0] + 1:
            problems.append(f"row {i}: accept cells not contiguous ({idx.tolist()})")
            continue
        a = "left" if idx[0] == 0 else ("right" if idx[-1] == ncols - 1 else None)
        if a is None:
            problems.append(f"row {i}: accept run {idx[0]}..{idx[-1]} anchored at neither grid edge")
        elif idx.size < ncols:  # full rows touch both edges and constrain nothing
            if anchor is None:
                anchor = a
            elif a != anchor:
                problems.append(f"row {i}: anchor flips {anchor}->{a}")
    diffs = np.diff(lens)
    if not ((diffs >= 0).all() or (diffs <= 0).all()):
        problems.append(f"accept run lengths not monotone across rows: {lens}")
    return problems


def check_workbook(path: str, sheet: str = "Cutoff Grids", verbose: bool = True) -> int:
    """Check every grid on *sheet*; returns the number of failing grids."""
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    grids = find_grids(ws)
    if verbose:
        print(f"found {len(grids)} grids on {sheet!r}")
    failures = 0
    for r0, c0, label, nrows, ncols in grids:
        vals, frontier = read_grid(ws, r0, c0, nrows, ncols)
        problems = check_grid(vals, frontier)
        failures += bool(problems)
        if verbose:
            counts = {k: int((vals == k).sum()) for k in ("A", "R", "—")}
            status = "FAIL" if problems else "OK "
            print(f"[{status}] R{r0}C{c0} ({label}): {nrows}x{ncols}, {counts}")
            for p in problems:
                print(f"       - {p}")
    if verbose:
        print(f"\n{failures} grid(s) with problems" if failures else "\nAll grids pass.")
    return failures


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", nargs="?", default="output/consolidated_risk_production.xlsx")
    ap.add_argument("--sheet", default="Cutoff Grids")
    args = ap.parse_args()
    return 1 if check_workbook(args.workbook, args.sheet) else 0


if __name__ == "__main__":
    sys.exit(main())
